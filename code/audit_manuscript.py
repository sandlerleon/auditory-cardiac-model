# -*- coding: utf-8 -*-
"""Check the built manuscript against the model and against itself."""
import io, json, os, re
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
DOC = r"C:\Users\Leon\Downloads\Medical Hypotheses\Auditory_Cardiac_Channel_MedicalHypotheses.docx"
R = json.load(io.open(os.path.join(HERE, "escalation_results.json"), encoding="utf-8"))

d = Document(DOC)
paras = [p.text for p in d.paragraphs]
tbl = []
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            tbl.append(c.text)
TEXT = "\n".join(paras + tbl)

CHECKS = [
 ("baseline crossing day", "%.2f" % R["baseline"]["cross_day"]),
 ("delay hours",           "%.1f" % R["intervention"]["delay_hours"]),
 ("ceiling hours",         "%.1f" % R["max_delay_hours"]),
 ("no-loop crossing day",  "%.2f" % R["no_affective_loop"]["cross_day"]),
 ("dose change pct",       "%.0f" % (-100 * R["intervention"]["mean_dose_change"])),
 ("sensitivity min delay", "%.1f" % min(x["delay_hours"] for x in R["sensitivity"])),
 ("sensitivity max ceiling", "%.1f" % max(x["ceiling_hours"] for x in R["sensitivity"])),
 ("parameter tau_CL",      "%g" % R["parameters"]["tau_CL"]),
]
print("NUMERIC CHECKS")
missing = 0
for name, val in CHECKS:
    ok = val in TEXT
    print("   %-24s %-8s %s" % (name, val, "ok" if ok else "NOT FOUND"))
    missing += 0 if ok else 1

# every reference cited, every citation listed
i = next(k for k, t in enumerate(paras) if t.strip() == "References")
body = "\n".join(paras[:i] + tbl)
reflist = [t for t in paras[i + 1:] if t.strip()]
cited = set()
for m in re.finditer(r"\[([\d,\s\u2013-]+)\]", body):
    for part in m.group(1).split(","):
        part = part.strip()
        if part.isdigit():
            cited.add(int(part))
listed = set(range(1, len(reflist) + 1))
print("\nREFERENCE CHECKS")
print("   entries            : %d" % len(reflist))
print("   never cited        : %s" % (sorted(listed - cited) or "none"))
print("   cited, not listed  : %s" % (sorted(cited - listed) or "none"))

# claims the paper must not make
print("\nSAFETY-LANGUAGE CHECKS")
FORBIDDEN = [
 (r"\breplaces? (?:the )?(?:morphine|opioid)", "claims to replace opioid"),
 (r"\beliminat\w+ (?:the )?(?:need for )?opioid", "claims to eliminate opioid need"),
 (r"\bprevents? (?:the )?(?:neurotox|crossing|threshold)", "claims prevention"),
 (r"\bshould be (?:given|administered|used) instead", "recommends substitution"),
]
issues = []
for pat, why in FORBIDDEN:
    for m in re.finditer(pat, TEXT, re.I):
        seg = TEXT[max(0, m.start() - 70):m.end() + 70].replace("\n", " ")
        denial = re.search(r"(never|not|no \w+ strength|cannot) prevent|no run did|no perturbation allows", seg, re.I)
        if denial or "never a target" in seg:
            continue
        issues.append("%s -> ...%s..." % (why, seg))
print("   %s" % ("none found" if not issues else "%d found" % len(issues)))
io.open(os.path.join(HERE, "_audit_issues.txt"), "w", encoding="utf-8").write(
    "\n\n".join(issues) if issues else "none")

required = ["Ethical statement", "How this proposal could cause harm",
            "adjunctive", "never a target", "treatment recommendation",
            "modulating opioid escalation", "illustrative model threshold",
            "Declaration of generative AI and AI-assisted technologies in the manuscript preparation process",
            "Feasibility of P3", "partially overlapping but non-equivalent"]
for bad in ["limiting opioid escalation", "overlaps almost completely",
            "days before death", "direct supporting evidence", "neurotoxicity threshold",
            "2016;2019", "2014;2019"]:
    if bad.lower() in TEXT.lower():
        issues.append("phrase the revision removed is back: %r" % bad)
print("\nREQUIRED CONTENT")
for r_ in required:
    print("   %-42s %s" % (r_, "present" if r_.lower() in TEXT.lower() else "MISSING"))

print("\nSUMMARY")
print("   verdict : %s" % ("PASS" if not (missing or issues
                                          or (listed - cited) or (cited - listed))
                           else "REVIEW NEEDED"))
