# -*- coding: utf-8 -*-
"""The separate items Medical Hypotheses' Editorial Manager requires alongside the
anonymised manuscript: title page, highlights, CRediT author statement, declaration
of interest statement and ethics statement.

The journal runs double-blind review, so everything that identifies the author
(name, affiliation, ORCID, e-mail, repository and archive links) lives on the title
page and nowhere in the manuscript file.

    python build_submission_items.py
"""
import io
import json
import os

from docx import Document
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\Leon\Downloads\Medical Hypotheses"
R = json.load(io.open(os.path.join(HERE, "escalation_results.json"), encoding="utf-8"))

TITLE = ("The auditory\u2013cardiac channel in the last days of life: a neurovisceral "
         "hypothesis for modulating opioid escalation")
AUTHOR = "Leon Sandler"
AFFIL = "Independent Researcher, Northbrook, Illinois 60062, United States"
EMAIL = "sandler.leon@gmail.com"
ORCID = "0009-0007-4584-808X"

HIGHLIGHTS = [
    "Terminal opioid titration may partly track autonomic arousal, not pain alone",
    "Hearing responses persist in dying patients; semantic processing is untested",
    "A minimal model shows damping arousal delays, never prevents, accumulation",
    "Five falsifiable predictions, including an HRV test with no treatment change",
    "Adjunctive only: analgesia is never withheld; dose reduction is never a target",
]
for h in HIGHLIGHTS:
    assert len(h) <= 85, (len(h), h)          # Elsevier: max 85 characters incl. spaces


def new_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.3
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1.0)
    d.core_properties.author = AUTHOR
    return d


def para(d, text, bold=False, size=11, italic=False, after=8):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    return p


def save(d, name):
    path = os.path.join(OUT, name)
    d.save(path)
    print("saved", path)


# ------------------------------------------------------------------ title page
d = new_doc()
para(d, "Title page", bold=True, size=12, after=14)
para(d, TITLE, bold=True, size=14, after=14)
para(d, "Author", bold=True, after=2)
para(d, "%s (ORCID %s)" % (AUTHOR, ORCID), after=2)
para(d, AFFIL, after=12)
para(d, "Corresponding author", bold=True, after=2)
para(d, "%s, %s. E-mail: %s" % (AUTHOR, AFFIL, EMAIL), after=12)
para(d, "Article type", bold=True, after=2)
para(d, "Hypothesis paper", after=12)
para(d, "Funding", bold=True, after=2)
para(d, "This work received no external funding.", after=12)
para(d, "Declaration of competing interest", bold=True, after=2)
para(d, "The author declares no competing interests.", after=12)
para(d, "Acknowledgements", bold=True, after=2)
para(d, "None.", after=12)
para(d, "Data availability (identifying links withheld from the blinded manuscript)",
     bold=True, after=2)
para(d, "The model implementation, parameter set and figure generators are openly "
        "available at https://github.com/sandlerleon/auditory-cardiac-model and "
        "archived at https://doi.org/10.5281/zenodo.22860430. A preprint of this "
        "manuscript is deposited at https://doi.org/10.5281/zenodo.22860432. Both are "
        "concept DOIs that resolve to the current version. No patient data were "
        "generated or analysed.", after=12)
para(d, "Declaration of generative AI and AI-assisted technologies in the manuscript "
        "preparation process", bold=True, after=2)
para(d, "During the preparation of this work the author used Claude (Anthropic) in "
        "order to assist with literature synthesis, implementation of the illustrative "
        "model, and drafting and editing of the text. After using this tool, the author "
        "reviewed and edited the content as needed, verified every cited reference "
        "against its source, and takes full responsibility for the content of the "
        "published article.", after=12)
save(d, "Title_Page_Sandler.docx")

# ------------------------------------------------------------------ highlights
d = new_doc()
para(d, "Highlights", bold=True, size=12, after=10)
for h in HIGHLIGHTS:
    p = d.add_paragraph(style="List Bullet")
    p.add_run(h).font.size = Pt(11)
d.core_properties.author = ""                  # highlights go to reviewers too
save(d, "Highlights.docx")

# ------------------------------------------------------------------ CRediT
d = new_doc()
para(d, "CRediT authorship contribution statement", bold=True, size=12, after=10)
para(d, "%s: Conceptualization, Methodology, Software, Formal analysis, "
        "Investigation, Visualization, Writing \u2013 original draft, Writing \u2013 "
        "review & editing." % AUTHOR)
para(d, "The author is the sole contributor to this work.", italic=True)
save(d, "CRediT_Author_Statement_Sandler.docx")

# ------------------------------------------------------------------ declaration of interest
d = new_doc()
para(d, "Declaration of interests", bold=True, size=12, after=10)
para(d, "\u2612 The author declares no known competing financial interests or "
        "personal relationships that could have appeared to influence the work "
        "reported in this paper.")
para(d, "\u2610 The author declares the following financial interests/personal "
        "relationships which may be considered as potential competing interests:")
para(d, "Manuscript: \u201c%s\u201d" % TITLE, size=10, italic=True)
para(d, "%s, %s" % (AUTHOR, AFFIL), size=10)
save(d, "Declaration_of_Interest_Sandler.docx")

# ------------------------------------------------------------------ ethics statement
d = new_doc()
para(d, "Ethics statement", bold=True, size=12, after=10)
para(d, "This is a theoretical hypothesis paper. No human participants or animals were "
        "studied, no patient data or samples were collected or analysed, and no "
        "intervention described in the manuscript has been tested. Ethical approval "
        "and informed consent were therefore not required.")
para(d, "The manuscript does not constitute clinical advice or a treatment "
        "recommendation, and it states explicitly that nothing in it should delay, "
        "reduce or substitute for established palliative care, including adequate "
        "analgesia. Any prospective evaluation of the predictions it proposes would "
        "require approval by a research ethics committee, with particular attention to "
        "consent by proxy for patients who cannot consent, the right to decline, and "
        "safety monitoring for under-treated pain (manuscript Sections 8 and 10).")
save(d, "Ethics_Statement_Sandler.docx")
