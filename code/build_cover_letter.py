# -*- coding: utf-8 -*-
"""Cover letter for the Medical Hypotheses submission.

Numbers come from the model output, as in the manuscript.

    python build_cover_letter.py
"""
import io
import json
import os

from docx import Document
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\Leon\Downloads\Medical Hypotheses"
R = json.load(io.open(os.path.join(HERE, "escalation_results.json"), encoding="utf-8"))
INT, BASE = R["intervention"], R["baseline"]

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(10)
st.paragraph_format.line_spacing = 1.15
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1.0)
    s.top_margin = s.bottom_margin = Inches(1.0)


def P(text, bold=False, size=11, after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p


P("Leon Sandler", size=10.5, after=0)
P("Independent Researcher", size=10.5, after=0)
P("Northbrook, Illinois 60062, United States", size=10.5, after=0)
P("sandler.leon@gmail.com | ORCID 0009-0007-4584-808X", size=10.5, after=14)
P("The Editors", size=10.5, after=0)
P("Medical Hypotheses", size=10.5, after=14)

P("Dear Editors,", after=10)

P("I am submitting \u201cThe auditory\u2013cardiac channel in the last days of "
  "life: a neurovisceral hypothesis for limiting opioid escalation\u201d for "
  "consideration as a hypothesis paper.")

P("The clinical observation behind it is ordinary and, I think, under-theorised. "
  "In the last days of life, opioid infusion is titrated against observed "
  "distress, but the signs being titrated against \u2014 restlessness, "
  "grimacing, tachycardia \u2014 are the same signs produced by opioid-induced "
  "neurotoxicity as metabolites accumulate in a patient whose clearance is "
  "failing. The bedside cannot distinguish too little drug from too much, and "
  "the humane default is to give more. The cost is that patients become "
  "unreachable days before they die.")

P("The hypothesis", bold=True, after=4)
P("I propose that a component of that escalation is driven by a self-reinforcing "
  "loop between autonomic arousal and perceived pain rather than by nociception "
  "alone, and that the loop is reachable through the subcortical auditory "
  "pathway, which remains functional when other sensory channels have failed. "
  "The paper sets out the anatomy, states the evidentiary standing of every link "
  "in a table, and analyses a minimal model of the loop.")

P("Why I think it is worth publishing", bold=True, after=4)
for txt in [
    "It makes the mechanism quantitative rather than gestural. The model says "
    "the intervention would postpone the neurotoxic threshold by roughly %.0f "
    "hours at modest strength, with a hard ceiling near %.0f hours, and would "
    "never prevent the crossing at any strength. A bounded, unflattering "
    "prediction is more useful than an open-ended claim, and it is what makes "
    "the hypothesis refutable."
    % (INT["delay_hours"], R["max_delay_hours"]),

    "It separates what is known from what is proposed. Auditory responses "
    "persist in dying patients and semantic processing survives unresponsiveness "
    "in other populations, but no study has shown semantic processing in "
    "actively dying patients. Popular accounts routinely conflate these. I have "
    "made the conflated claim the paper's primary prediction instead of its "
    "premise.",

    "It states how it could cause harm. The practical implication of this "
    "hypothesis is a smaller opioid dose in patients who cannot report pain, "
    "which is a dangerous thing to propose casually. Section 8 enumerates the "
    "failure modes, and the paper is explicit that the proposal is adjunctive "
    "and that dose reduction is an outcome to be observed, never a target.",

    "The decisive early test is cheap and non-invasive. Whether structured "
    "verbal input measurably raises high-frequency heart rate variability in "
    "unresponsive dying patients requires no protocol deviation and would be "
    "informative whichever way it resolved.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(txt)
    r.font.size = Pt(11)

P("Openness", bold=True, after=4)
P("The model implementation, parameter set and figure generators are openly "
  "available at https://github.com/sandlerleon/auditory-cardiac-model and "
  "archived at https://doi.org/10.5281/zenodo.22860430; the manuscript is "
  "deposited at https://doi.org/10.5281/zenodo.22860432. Every number in the "
  "paper is read programmatically from the model output rather than "
  "transcribed, and all 17 references were verified against Crossref.")

P("Declarations", bold=True, after=4)
P("This manuscript is original, is not under consideration elsewhere, and has "
  "not been published previously other than as the archived preprint noted "
  "above. I am the sole author, I have no competing interests, and the work "
  "received no external funding. It is entirely theoretical: no patient was "
  "studied and no data were collected. Generative AI (Claude, Anthropic) was "
  "used to assist with literature synthesis, model implementation and drafting; "
  "I reviewed and edited all content, verified every reference, and take full "
  "responsibility for the manuscript, as stated in the Declarations.")

P("Thank you for considering it.", after=14)
P("Yours sincerely,", after=0)
P("Leon Sandler", after=0)

if not os.path.isdir(OUT):
    os.makedirs(OUT)
path = os.path.join(OUT, "MedicalHypotheses_Cover_Letter_Sandler.docx")
doc.save(path)
print("words : %d" % sum(len(p.text.split()) for p in doc.paragraphs))
print("saved : %s" % path)
