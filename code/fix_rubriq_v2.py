# -*- coding: utf-8 -*-
"""Repair the Rubriq-edited manuscript (…_v2.docx) without discarding its style edits.

Rubriq's pass converted the text to US English and tidied punctuation, which is kept.
It also changed the sense of about twenty passages, deleted the 'Established:' labels
from the evidence hierarchy, inverted the Figure 2 caption, and altered a reference
title. Each of those is restored here; every replacement is asserted to hit exactly
once, so nothing is fixed silently or twice.

Then writes the double-blind copy (no name, affiliation, ORCID, e-mail, repository or
archive links, blank document metadata).

    python fix_rubriq_v2.py
"""
import copy
import os
import shutil

from docx import Document
from docx.oxml.ns import qn

D = r"C:\Users\Leon\Downloads\Medical Hypotheses"
SRC = os.path.join(D, "Auditory_Cardiac_Channel_MedicalHypotheses_v2.docx")
RAW = os.path.join(D, "_rubriq_raw", "Auditory_Cardiac_Channel_MedicalHypotheses_v2_rubriq_raw.docx")
OUT = SRC
OUT_BLIND = os.path.join(D, "Auditory_Cardiac_Channel_MedicalHypotheses_v2_anonymized.docx")

FIXES = [
    # guideline: one guideline, and 'assessed' need is the point
    ("The NICE guidelines on the care of dying adults recommend",
     "The NICE guideline on the care of dying adults recommends"),
    ("according to individual needs [1]", "according to individually assessed need [1]"),
    ("tachypnoea", "tachypnea"),                                  # consistent US spelling
    # M3G: the mechanism is contested (the paper says so), not merely unclear
    ("the mechanism usually invoked for this neurotoxicity is unclear",
     "the mechanism usually invoked for this neurotoxicity is contested"),
    ("rather than as a specific mediator", "rather than asserting a specific mediator"),
    # 'last days' of life, not 'past days'
    ("escalation in the past days", "escalation in the last days"),
    ("It is not measured in dying patients", "It has not been measured in dying patients"),
    ("Afferent trafficking from baroreceptors", "Afferent traffic from baroreceptors"),
    # garbled sentence
    ("To determine whether the proposed loop can produce the clinical pattern and what "
     "degree of damping it would be worth doing so, we analyze",
     "To determine whether the proposed loop can produce the clinical pattern, and what "
     "damping it would be worth, we analyze"),
    # Figure 2 caption: meaning had been inverted
    ("and the lack of damping strength prevents crossing",
     "and no damping strength prevents the crossing"),
    ("The parameters are illustrated throughout", "The parameters are illustrative throughout"),
    ("If an intervention was shown to postpone", "If an intervention were shown to postpone"),
    ("particularly if responsiveness was preserved", "particularly if responsiveness were preserved"),
    # Section 6 opener lost words
    ("chains together claims very different evidentiary standing",
     "chains together claims of very different evidentiary standing"),
    ("assigns each status rather than folding it into",
     "assigns each a status rather than folding them into"),
    ("from conscious people who the argument does not depend on",
     "from conscious people that the argument does not depend on"),
    ("P3. Auditory input moves to the autonomic state.", "P3. Auditory input moves autonomic state."),
    ("the rest of the chain for opioid escalation", "the rest of the chain to opioid escalation"),
    # restore the commas that keep 'opioids and sedatives' out of the 'given for' phrase
    ("Anticholinergics given for respiratory secretions and opioids and sedatives themselves "
     "alter cardiac vagal activity",
     "Anticholinergics given for respiratory secretions, and opioids and sedatives "
     "themselves, alter cardiac vagal activity"),
    ("most likely to cause harm, and P4 specifies distress scores",
     "most likely to cause harm and is the reason P4 specifies distress scores"),
    # an advance (directive) preference governs; it is not 'governed'
    ("advanced preference should be governed", "advance preference should govern"),
    ("It is that a specific, measurable", "It is that a specific, measurable"),  # placeholder, see below
    ("and no tool is used to produce or alter data", "and none was used to produce or alter data"),
    # reference title altered by the copyedit
    ("Anterior Cingulate However, Not Somatosensory Cortex",
     "Anterior Cingulate But Not Somatosensory Cortex"),
]
FIXES = [f for f in FIXES if f[0] != f[1]]
# Section 9: 'Specifically, a specific, measurable…; the relevant…; and the mechanism…'
FIXES.append(("Specifically, a specific, measurable autonomic pathway mediates the effect; "
              "the relevant outcome", "It is that a specific, measurable autonomic pathway "
              "mediates the effect; that the relevant outcome"))
FIXES.append(("comfort score; and the mechanism predicts",
              "comfort score; and that the mechanism predicts"))


def all_paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def replace_in_paragraph(p, old, new):
    """Replace text that may span several runs; formatting of the first run is kept."""
    runs = p.runs
    text = "".join(r.text for r in runs)
    i = text.find(old)
    if i < 0:
        return False
    j = i + len(old)
    pos, first = 0, None
    for r in runs:
        a, b = pos, pos + len(r.text)
        pos = b
        if b <= i or a >= j:
            continue
        s, e = max(i, a) - a, min(j, b) - a
        if first is None:
            r.text = r.text[:s] + new + r.text[e:]
            first = r
        else:
            r.text = r.text[:s] + r.text[e:]
    return True


def apply_fixes(doc):
    seen = set()
    for old, new in FIXES:
        hits = 0
        for p in all_paragraphs(doc):
            if id(p._p) in seen and False:
                continue
            if old in p.text:
                assert replace_in_paragraph(p, old, new)
                hits += 1
        assert hits == 1, "%r hit %d times" % (old[:60], hits)
    # the evidence hierarchy: restore the two deleted 'Established:' labels
    for start in ("Auditory information is processed in brainstem",
                  "Autonomic regulation involves brainstem"):
        p = next(q for q in doc.paragraphs if q.text.startswith(start))
        lead = copy.deepcopy(p.runs[0]._r)
        p.runs[0]._r.addprevious(lead)
        from docx.text.run import Run
        r = Run(lead, p)
        r.text = "Established: "
        r.bold = True
    p = next(q for q in doc.paragraphs if q.text.startswith("Untested: whether"))
    assert replace_in_paragraph(p, "Untested: whether", "Untested: Whether")
    # the bold lead was a separate run; re-bold only the label
    for r in p.runs:
        if r.text.startswith("Untested: Whether"):
            r.text = "Untested: "
            tail = copy.deepcopy(r._r)
            r._r.addnext(tail)
            from docx.text.run import Run
            t = Run(tail, p)
            t.text = "Whether"
            t.bold = None
            break


def make_blind(doc):
    body = doc.element.body
    for start in ("Leon Sandler", "Independent Researcher, Northbrook", "ORCID: 0009"):
        p = next(q for q in doc.paragraphs if q.text.startswith(start))
        body.remove(p._p)
    p = next(q for q in doc.paragraphs if q.text.startswith("Data availability."))
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = (
        "Data availability. The model implementation, parameter set and figure generators "
        "are openly available in a public repository with an archived, versioned DOI; the "
        "links are withheld here for double-blind review and are given on the title page. "
        "Running escalation_model.py followed by make_figures.py reproduces every number "
        "and figure. The equations, initial conditions and parameter values needed to "
        "reproduce the model without the code are given in Section 5 and Tables 1\u20132. "
        "No patient data were generated or analyzed.")
    cp = doc.core_properties
    cp.author = cp.last_modified_by = cp.comments = ""
    text = "\n".join(q.text for q in all_paragraphs(doc)).lower()
    for ident in ("sandler", "leon", "0009-0007", "sandlerleon", "zenodo", "northbrook",
                  "gmail", "github"):
        assert ident not in text, ident


def main():
    os.makedirs(os.path.dirname(RAW), exist_ok=True)
    if not os.path.exists(RAW):
        shutil.copy2(SRC, RAW)                 # keep Rubriq's untouched output
    doc = Document(RAW)
    apply_fixes(doc)
    doc.core_properties.author = "Leon Sandler"
    doc.save(OUT)
    print("repaired  :", OUT)
    blind = Document(OUT)
    make_blind(blind)
    blind.save(OUT_BLIND)
    print("anonymized:", OUT_BLIND)


if __name__ == "__main__":
    main()
