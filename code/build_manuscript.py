# -*- coding: utf-8 -*-
"""Build the Medical Hypotheses manuscript.

Every number is read from escalation_results.json, and every reference from the
Crossref verification pass, so neither can drift from its source.

Three things the source draft asserted that the literature does not support have
been reworked rather than repeated:

  * N400 semantic responses have NOT been recorded in actively dying patients.
    Tone-based auditory ERPs have (Blundon 2020); N400 has been recorded in
    disorders of consciousness (Steppacher 2013). The semantic claim in dying
    patients is this paper's primary testable prediction, not its premise.
  * M3G neurotoxicity is contested, and the refutation comes from the same group
    that proposed it.
  * Music therapy and hypnotherapy at the end of life already exist and have
    been reviewed; the novelty claim has to be narrower than the draft's.

Added because the subject demands it: an explicit harm section and an ethical
statement. The draft had neither, in a paper about reducing opioid dose in
dying patients.

    python build_manuscript.py
"""
import io
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\Leon\Downloads\Medical Hypotheses"
R = json.load(io.open(os.path.join(HERE, "escalation_results.json"), encoding="utf-8"))
REFS = json.load(io.open(os.path.join(HERE, "_refs_final.json"), encoding="utf-8"))

BASE = R["baseline"]
INT = R["intervention"]
P = R["parameters"]

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 2.0

for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1.0)
    s.top_margin = s.bottom_margin = Inches(1.0)
    sectPr = s._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:start"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    anchor = sectPr.find(qn("w:pgMar"))
    if anchor is None:
        anchor = sectPr.find(qn("w:pgSz"))
    if anchor is None:
        sectPr.insert(0, ln)
    else:
        anchor.addnext(ln)

# citation numbering, by order of first appearance
ORDER = []


def C(*tags):
    nums = []
    for t in tags:
        if t not in ORDER:
            ORDER.append(t)
        nums.append(ORDER.index(t) + 1)
    nums.sort()
    return "[%s]" % ",".join(str(n) for n in nums)


def H(text, size=12, before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    return p


def Pp(text, indent=False, size=11, spacing=2.0, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = spacing
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.3)
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    return p


def BULLET(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.5
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.size = Pt(11)
    r = p.add_run(text)
    r.font.size = Pt(11)
    return p


def CAP(text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    return p


def FIG(name, width=6.2):
    path = os.path.join(HERE, name)
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def TBL(headers, rows, fs=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].paragraph_format.line_spacing = 1.0
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(fs)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            pp = cells[i].paragraphs[0]
            pp.paragraph_format.line_spacing = 1.0
            pp.paragraph_format.space_after = Pt(1)
            pp.add_run(str(v)).font.size = Pt(fs)
    return t


# ==================================================================== FRONT
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.5
TITLE = ("The auditory\u2013cardiac channel in the last days of life: a "
         "neurovisceral hypothesis for modulating opioid escalation")
r = p.add_run(TITLE)
r.bold = True
r.font.size = Pt(14)

Pp("Leon Sandler", spacing=1.5)
Pp("Independent Researcher, Northbrook, Illinois, United States", spacing=1.5, size=10.5)
Pp("ORCID: 0009-0007-4584-808X \u00b7 Correspondence: sandler.leon@gmail.com",
   spacing=1.5, size=10.5)

H("Abstract")
ABSTRACT = (
    "In the last days of life, opioid infusion is commonly titrated against "
    "observed signs of distress, but these signs may reflect nociception, "
    "autonomic arousal, delirium, or opioid-related neurotoxicity. Where "
    "clearance of active metabolites declines, escalation and accumulation can "
    "interact: accumulation may produce agitation, agitation may be read as "
    "pain, and perceived pain may prompt further escalation. We propose that "
    "part of this escalation is driven not by nociception but by a "
    "self-reinforcing loop between autonomic arousal and perceived distress, "
    "and that the loop may be accessible through hearing. We separate what is "
    "established (auditory information is processed in brainstem circuits "
    "before it reaches cortex; autonomic regulation involves brainstem, "
    "limbic and cortical structures; tone-evoked auditory potentials persist "
    "in actively dying patients) from what is hypothesised (that residual "
    "auditory processing can be used to lower autonomic arousal in dying "
    "patients) and from what is untested (whether this changes opioid "
    "requirement). A minimal, fully specified model of the loop, with "
    "illustrative parameters and a sensitivity analysis, indicates that "
    "damping arousal postpones rather than prevents metabolite accumulation, "
    "with a benefit bounded at about a day in the base case. We give five "
    "falsifiable predictions, a "
    "feasibility protocol for heart rate variability measurement in dying "
    "patients, and the conditions under which the proposal would cause harm. "
    "No patient data were collected, and nothing here is a treatment "
    "recommendation or a reason to withhold or delay analgesia.")
Pp(ABSTRACT, spacing=1.5)

Pp("Keywords: palliative care; opioid escalation; terminal agitation; auditory "
   "event-related potentials; neurovisceral integration; heart rate variability",
   spacing=1.5, size=10.5)

# ================================================================= SECTION 1
H("1. The problem")
Pp("Care of the dying is, pharmacologically, a titration problem. Continuous "
   "subcutaneous or intravenous opioid is adjusted against observed signs of "
   "distress: grimacing, restlessness, tachypnoea, tachycardia. The NICE "
   "guideline on care of dying adults recommends that medicines for symptoms "
   "including pain be prescribed and adjusted according to individual assessed "
   "need %s (summarised in %s), and opioids are the best-evidenced treatment "
   "for refractory breathlessness in advanced disease %s. None of this is in "
   "question here."
   % (C("nice_ng31"), C("nice2016"), C("barnes2016")))
Pp("Throughout this paper, pain, distress and agitation are treated as partially "
   "overlapping but non-equivalent clinical constructs. The hypothesis concerns "
   "the contribution of autonomic arousal to observed distress and to opioid "
   "titration; it does not claim that terminal agitation is itself pain.",
   indent=True)
Pp("The difficulty is that the signs being titrated against are not specific to "
   "nociception, and the patient can no longer disambiguate them. Where renal "
   "function declines, as it may in the last days of life, active opioid "
   "metabolites can accumulate. Opioid-induced neurotoxicity \u2014 myoclonus, "
   "hyperalgesia, delirium, agitation \u2014 is a recognised complication in "
   "palliative medicine %s, and some of its features, such as restlessness and "
   "apparent distress, can resemble under-treated pain. A clinician watching a "
   "restless, grimacing patient may not be able to tell from bedside signs "
   "alone whether the dose is too low or too high, and a common, humane "
   "response is to increase it."
   % C("oin2026"), indent=True)
Pp("We note at once that the mechanism usually invoked for this neurotoxicity is "
   "contested. Morphine-3-glucuronide was proposed as the neuroexcitatory agent "
   "%s, but the same group subsequently reported that M3G is not neurotoxic to "
   "primary hippocampal or cerebellar granule neurones %s. The clinical "
   "syndrome is well described; its molecular basis is not settled, and we "
   "therefore treat metabolite accumulation as a phenomenological driver rather "
   "than asserting a specific mediator."
   % (C("smith2000"), C("m3g_refute")), indent=True)
Pp("When escalation does occur, one potential cost is reduced capacity for "
   "communication, since both sedation and delirium limit interaction with "
   "those at the bedside. How often this happens, and how early, is not "
   "established here and is not assumed; the hypothesis concerns a mechanism "
   "that could contribute to it, and preserved responsiveness is the outcome "
   "the proposal ultimately seeks to measure (prediction P5).", indent=True)

# ================================================================= SECTION 2
H("2. The hypothesis")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.4)
p.paragraph_format.right_indent = Inches(0.4)
p.paragraph_format.line_spacing = 1.5
r = p.add_run("A component of terminal opioid escalation is driven by a "
              "self-reinforcing loop between autonomic arousal and perceived "
              "distress, rather than by nociceptive input alone. Because "
              "auditory processing appears to persist in dying patients when "
              "other channels have failed, structured auditory input may be "
              "able to lower autonomic arousal and damp that loop, postponing "
              "the point at which accumulation forces a choice between comfort "
              "and consciousness.")
r.bold = True
r.font.size = Pt(11)
Pp("Two claims are separable here and should be judged separately. The first is "
   "that an arousal\u2013distress loop contributes materially to dose escalation "
   "in the last days. The second is that the auditory channel is a usable point "
   "of entry to it. The first could be true and the second false. The chain "
   "that links them has four levels of evidence, and neuroanatomical "
   "connectivity at the first level is not evidence of a therapeutic effect at "
   "the fourth:", indent=True)
BULLET("auditory information is processed in brainstem circuits before it "
       "reaches cortex, and some of those circuits drive fast responses, such as "
       "the acoustic startle, without cortical involvement %s." % C("koch1999"),
       bold_lead="Established: ")
BULLET("autonomic regulation involves brainstem, limbic and cortical structures "
       "%s." % C("benarroch1993", "thayer2009"), bold_lead="Established: ")
BULLET("structured auditory stimulation in actively dying patients can exploit "
       "residual auditory processing to alter autonomic state.",
       bold_lead="Hypothesis: ")
BULLET("whether this produces a clinically meaningful change in opioid "
       "escalation.", bold_lead="Untested: ")
Pp("Neither claim implies that analgesia should be withheld, reduced or delayed. "
   "Throughout, the opioid is titrated to comfort exactly as it would otherwise "
   "be; any change in requirement is an outcome to be observed, never a target "
   "to be pursued (Section 8).", indent=True)

# ================================================================= SECTION 3
H("3. Why the auditory channel, and what has actually been shown")
Pp("Acoustic signals reach the brainstem through the cochlear nuclei, superior "
   "olivary complex and inferior colliculus before thalamocortical routing, and "
   "some brainstem auditory circuits act without cortical involvement: the "
   "acoustic startle response is mediated by a circuit in the lower brainstem "
   "%s. In animal studies of fear conditioning, the amygdala receives auditory "
   "input directly from the auditory thalamus as well as from auditory cortex, "
   "and conditioned fear responses to simple tones can be acquired through the "
   "thalamic route %s. In healthy people, functional neuroimaging shows that "
   "music modulates activity in structures involved in autonomic regulation, "
   "including the amygdala, hypothalamus, insula and cingulate cortex %s. "
   "Whether any of these routes remains functional, or usable, in actively "
   "dying patients is not known."
   % (C("koch1999"), C("ledoux2000"), C("koelsch2014")))
Pp("The evidentiary position must be stated precisely, because the source "
   "literature is often summarised more strongly than it supports. Two distinct "
   "findings are relevant and they are not the same finding.", indent=True)
BULLET("Auditory event-related potentials to tone sequences persist in actively "
       "dying hospice patients, in some cases within hours of death, and "
       "resemble those of healthy controls %s. This establishes that the "
       "auditory system is still processing structured acoustic input at the end "
       "of life. It does not establish that words are being understood."
       % C("blundon2020"), bold_lead="Hearing persists. ")
BULLET("The N400, an index of semantic mismatch, has been recorded in patients "
       "with disorders of consciousness, where its presence predicts later "
       "recovery %s. This establishes that automatic semantic processing can "
       "survive profound unresponsiveness. It was not measured in dying "
       "patients, and its absence in many such patients is precisely what makes "
       "it prognostic." % C("steppacher2013"),
       bold_lead="Semantic processing can survive unresponsiveness. ")
Pp("No published study has demonstrated semantic processing in actively dying "
   "patients. The proposition that the dying brain decodes the meaning of what "
   "is said to it is therefore a hypothesis, not a premise, and we treat it as "
   "this paper's primary experimental prediction (Section 7). Framing it as "
   "established \u2014 as popular accounts of \u2018hearing is the last sense to "
   "go\u2019 routinely do \u2014 would put the weight of the argument on a "
   "claim the literature has not yet tested.", indent=True)

# ================================================================= SECTION 4
H("4. From acoustic input to autonomic state")
Pp("The route from processed sound to cardiac output runs through the central "
   "autonomic network, a set of brainstem, hypothalamic, limbic and cortical "
   "structures that regulate autonomic outflow %s. Under the neurovisceral "
   "integration account, prefrontal and limbic structures exert inhibitory "
   "control over sympathoexcitatory circuits, with vagally mediated heart rate "
   "variability as the accessible index of that control %s. Afferent traffic "
   "from baroreceptors returns through the nucleus tractus solitarius to the "
   "insular cortex, which supports interoceptive representation of bodily "
   "state %s. In healthy people the loop is bidirectional: cardiac state is an "
   "input to affective state as well as an output of it. None of these links "
   "has been examined in actively dying patients."
   % (C("benarroch1993"), C("thayer2009"), C("critchley2004")))
Pp("Hypnosis offers a supporting analogy rather than direct evidence. In "
   "healthy volunteers, suggestion directed at unpleasantness rather than "
   "intensity altered anterior cingulate activity without corresponding change "
   "in primary somatosensory cortex %s, showing that the affective component "
   "of pain can be moved while the sensory component is left intact, and "
   "meta-analysis supports an analgesic effect of hypnotic procedures %s. These "
   "findings come from conscious, responsive people and are not evidence that "
   "the same effect occurs in dying patients. The argument does not depend on "
   "them; the chain it does depend on is auditory processing \u2192 autonomic "
   "state \u2192 distress \u2192 opioid titration \u2192 metabolite accumulation."
   % (C("rainville1997"), C("montgomery2000")), indent=True)
FIG("figure1_pathway.png")
CAP("Figure 1. The proposed auditory\u2013cardiac pathway. Grey boxes are "
    "established anatomy. Green marks the only step measured in dying patients: "
    "tone-evoked auditory event-related potentials %s; semantic processing has "
    "not been measured in this population. Orange dashed boxes and arrows are "
    "proposed links that are untested in dying patients; the subcortical route "
    "to the amygdala is described in animal studies %s. Blue marks the proposed "
    "intervention and the proposed readout. The red dashed limb is the proposed "
    "feedback from interoceptive state to threat appraisal. Anatomical "
    "connection is not evidence of a therapeutic effect."
    % (C("blundon2020"), C("ledoux2000")))

# ================================================================= SECTION 5
H("5. A minimal model of the escalation loop")
Pp("To ask whether the proposed loop can produce the clinical pattern, and what "
   "damping it would be worth, we analyse a deliberately minimal system. "
   "Perceived pain P is nociceptive drive N amplified by autonomic arousal A "
   "with gain g; arousal is driven by perceived pain and damped by a term V "
   "standing for the intervention; dose D is titrated toward perceived pain up "
   "to a ceiling; metabolite M accumulates as clearance declines "
   "exponentially; and once M first exceeds an illustrative model threshold M* "
   "an agitation term k_agit is added to N, which is what closes the loop. M* "
   "is a model construct, not an estimate of any neurotoxic concentration.")
for eq in ["P = N(1 + gA),   N = N\u2080 + k_agit\u00b7H(M \u2212 M*)",
           "dA/dt = (a\u2080 + k_P\u00b7P \u2212 V \u2212 A)/\u03c4_A,   A \u2265 0",
           "dD/dt = (min(P, D_max) \u2212 D)/\u03c4_D",
           "dM/dt = k_M\u00b7D \u2212 CL(t)\u00b7M,   CL(t) = CL\u2080\u00b7exp(\u2212t/\u03c4_CL)"]:
    Pp(eq, italic=True, spacing=1.2)
Pp("H is the unit step function. Initial conditions are A(0) = a\u2080, D(0) = 0 "
   "and M(0) = 0. The system is integrated by forward Euler with a step of %.3f "
   "day over %.0f days, and the crossing time is the first time at which M "
   "exceeds M*. Parameter values and units are listed in Table 1; with them, "
   "the equations above reproduce every number reported in this section."
   % (R["dt"], R["t_end"]), indent=True)
Pp("Every parameter is illustrative. None is fitted to patient data, this is not "
   "a pharmacokinetic simulation of morphine, and no quantity below should be "
   "read as a predicted dose or survival time. The model is used only to ask "
   "what class of behaviour the loop can produce and what would have to be "
   "measured to refute it. Results are therefore reported as relative timings.",
   indent=True)
PARAM_ROWS = [
    ("N\u2080", "N0", "baseline nociceptive drive", "a.u."),
    ("g", "g", "affective gain: amplification of perceived pain by arousal", "per unit arousal"),
    ("a\u2080", "a0", "baseline autonomic arousal", "a.u."),
    ("k_P", "kP", "drive of arousal by perceived pain", "arousal per unit pain"),
    ("\u03c4_A", "tau_A", "arousal time constant", "day"),
    ("\u03c4_D", "tau_D", "dose-adjustment time constant", "day"),
    ("k_M", "kM", "metabolite formation per unit dose", "M units per dose unit"),
    ("CL\u2080", "CL0", "initial metabolite clearance", "per day"),
    ("\u03c4_CL", "tau_CL", "clearance decay time constant", "day"),
    ("M*", "Mstar", "illustrative model threshold", "M units"),
    ("k_agit", "kAgit", "agitation added to N after M first exceeds M*", "a.u."),
    ("D_max", "Dmax", "dose ceiling", "dose units"),
]
TBL(["Symbol", "Value", "Meaning", "Units"],
    [[s, "%g" % P[k], m, u] for s, k, m, u in PARAM_ROWS]
    + [["V", "0; 0.35; 1.5; swept 0\u20133", "vagal damping (no intervention; "
        "modest; maximal in Fig. 2; sweep in Fig. 2B)", "arousal units"]])
CAP("Table 1. Model parameters, all illustrative and none fitted to patient data. "
    "a.u., arbitrary units. Integration: forward Euler, step %.3f day, %.0f days."
    % (R["dt"], R["t_end"]))
Pp("Three results follow. First, the loop is capable of producing the clinical "
   "pattern: with the assumed parameters the threshold is crossed on day "
   "%.2f, after which dose saturates. Second, damping arousal postpones that "
   "crossing \u2014 a modest damping term delays it by %.1f hours and lowers "
   "mean pre-threshold dose by %.0f%%. Third, and most importantly, damping "
   "never prevents the crossing at any strength. The delay saturates at a "
   "ceiling of %.1f hours, reached when arousal is fully suppressed; removing "
   "the affective loop altogether gives essentially the same answer (day %.2f). "
   "The benefit this hypothesis can claim is bounded, and the bound is hours to "
   "a little over a day."
   % (BASE["cross_day"], INT["delay_hours"], -100 * INT["mean_dose_change"],
      R["max_delay_hours"], R["no_affective_loop"]["cross_day"]), indent=True)
SENS = R["sensitivity"]
_ok = [s for s in SENS if s["baseline_cross_day"] is not None]
_pre = [s for s in SENS if s["param"] not in ("kAgit", "Dmax")]
Pp("These results are conditional on the assumed parameters, so each of the "
   "%d parameters in Table 1 was scaled in turn by 0.75 and 1.25 (Table 2). "
   "Across the %d runs the baseline crossing moved between day %.2f and day "
   "%.2f, the delay from modest damping between %.1f and %.1f hours, and the "
   "ceiling between %.1f and %.1f hours. In %s did any damping strength prevent "
   "the crossing. The agitation term and the dose ceiling act only after the "
   "crossing and leave all three quantities unchanged. The qualitative result "
   "— a delay that is real, bounded at hours to about two days, and never "
   "a prevention — is therefore not an artefact of one parameter choice, "
   "though its size is."
   % (len(set(s["param"] for s in SENS)), len(SENS),
      min(s["baseline_cross_day"] for s in _ok), max(s["baseline_cross_day"] for s in _ok),
      min(s["delay_hours"] for s in _ok), max(s["delay_hours"] for s in _ok),
      min(s["ceiling_hours"] for s in _ok), max(s["ceiling_hours"] for s in _ok),
      "no run" if not any(s["prevented"] for s in SENS) else "some runs"), indent=True)
_by = {}
for s in SENS:
    _by.setdefault(s["param"], {})[s["factor"]] = s
_sym = {k: sym for sym, k, _, _ in PARAM_ROWS}
TBL(["Parameter", "×0.75: crossing (day)", "delay (h)", "ceiling (h)",
     "×1.25: crossing (day)", "delay (h)", "ceiling (h)"],
    [[_sym[k]] + ["%.2f" % _by[k][f]["baseline_cross_day"] if j == 0 else
                  "%.1f" % _by[k][f][("delay_hours", "ceiling_hours")[j - 1]]
                  for f in (0.75, 1.25) for j in range(3)]
     for k in _by], fs=8.5)
CAP("Table 2. One-at-a-time sensitivity. Each parameter is scaled by 0.75 or 1.25 "
    "with all others at the Table 1 values; the baseline case is crossing day "
    "%.2f, delay %.1f h at V = 0.35, ceiling %.1f h. No perturbation allows any "
    "damping strength (V swept 0–3) to prevent the crossing."
    % (BASE["cross_day"], INT["delay_hours"], R["max_delay_hours"]))
FIG("figure2_model.png")
CAP("Figure 2. Behaviour of the minimal loop model. (A) Accumulated metabolite "
    "against time for no intervention, modest damping and maximal damping; the "
    "dotted line is the illustrative model threshold M*, not a neurotoxic "
    "concentration, and circles mark its crossing. (B) Delay in reaching the "
    "threshold as a function of damping strength. The delay saturates at %.1f "
    "hours and no damping strength prevents the crossing. Parameters are "
    "illustrative throughout (Table 1); the figure shows the shape of the "
    "result, not a prediction for any patient." % R["max_delay_hours"])
Pp("That the modelled benefit is bounded is not a weakness of the proposal but "
   "the most useful thing the model says. If an intervention were shown to "
   "postpone accumulation-driven agitation without compromising comfort, the "
   "resulting interval could be clinically meaningful, particularly if "
   "responsiveness were preserved. An intervention sold as a way to avoid "
   "opioid escalation would be both wrong and dangerous.", indent=True)

# ================================================================= SECTION 6
H("6. Evidence status of each link")
Pp("The argument above chains together claims of very different evidentiary "
   "standing. Table 3 assigns each a status rather than folding them into a "
   "single judgement about the hypothesis, so that a reader can locate the weak "
   "links directly.")
TBL(["Link", "Status", "Basis"],
    [["Brainstem auditory circuits can drive responses without cortical "
      "involvement", "Demonstrated", "Acoustic startle circuit in the lower "
      "brainstem " + C("koch1999")],
     ["A direct auditory thalamus → amygdala route supports conditioned fear "
      "responses to tones", "Demonstrated (animals)", "Fear-conditioning circuitry "
      + C("ledoux2000")],
     ["Sound modulates activity in limbic structures involved in autonomic "
      "regulation", "Demonstrated (healthy humans)", "Functional neuroimaging of "
      "music-evoked emotion " + C("koelsch2014")],
     ["Auditory ERPs to structured tone sequences persist in actively dying patients",
      "Demonstrated", "Direct measurement in hospice patients " + C("blundon2020")],
     ["Automatic semantic processing can survive profound unresponsiveness",
      "Demonstrated", "N400 in disorders of consciousness " + C("steppacher2013")],
     ["Semantic processing persists in actively dying patients",
      "Untested", "Not measured in this population; prediction P1"],
     ["Vagally mediated HRV indexes central autonomic control",
      "Demonstrated (healthy humans)", "Neurovisceral integration " + C("thayer2009")],
     ["Hypnotic suggestion separates pain affect from pain sensation",
      "Analogy only", "ACC dissociation in healthy volunteers " + C("rainville1997")],
     ["Hypnotic procedures reduce analgesic requirement",
      "Analogy only", "Randomised evidence in conscious sedation " + C("faymonville1997")
      + "; meta-analysis " + C("montgomery2000")],
     ["Central mechanisms can amplify pain beyond the peripheral input "
      "(central sensitisation)", "Demonstrated (other populations)",
      "Mechanistic review " + C("woolf2011") + "; a contribution from arousal is "
      "part of this proposal, not of the review"],
     ["An arousal–distress loop drives a material share of terminal escalation",
      "Proposed", "This paper; prediction P2"],
     ["Structured auditory input alters autonomic state in dying patients",
      "Untested", "This paper; prediction P3"],
     ["That change alters opioid escalation",
      "Untested", "This paper; prediction P4"],
     ["Damping postpones rather than prevents metabolite accumulation",
      "Model result", "Section 5; conditional on illustrative parameters (Tables 1–2)"],
     ["A specific metabolite mediates opioid-induced neurotoxicity",
      "Contested", "Proposed " + C("smith2000") + ", refuted by the same group "
      + C("m3g_refute")]])
CAP("Table 3. Status of each link in the chain. Rows 6, 11, 12 and 13 are the "
    "load-bearing untested claims and are the targets of the experimental "
    "programme in Section 7. 'Analogy only' marks evidence from conscious people "
    "that the argument does not depend on.")

# ================================================================= SECTION 7
H("7. Predictions and falsification")
Pp("The hypothesis is refuted by any of the following.")
for tag, body in [
 ("P1. Semantic processing in dying patients. ",
  "N400 responses to semantic mismatch, recorded in actively dying patients "
  "using the paradigm already validated in disorders of consciousness, should "
  "be present in a substantial minority. If they are absent in essentially all "
  "patients while tone-based responses persist, the verbal content of the "
  "intervention is irrelevant and only prosody and familiarity could matter."),
 ("P2. An arousal component to escalation. ",
  "Among patients matched for disease, site and nociceptive burden, those with "
  "lower vagally mediated heart rate variability should require larger opioid "
  "dose increments over the same interval. If dose trajectory is independent of "
  "autonomic state, the loop does not contribute materially and the proposal "
  "fails at its foundation."),
 ("P3. Auditory input moves autonomic state. ",
  "Structured verbal pacing delivered to unresponsive dying patients should "
  "produce a measurable increase in high-frequency heart rate variability "
  "relative to matched acoustic control, within minutes. This is the most "
  "directly testable prediction and requires no change to any treatment."),
 ("P4. Dose trajectory. ",
  "In a randomised comparison against an attention-matched control, the "
  "intervention arm should show a lower rate of dose escalation over 48 hours "
  "without any increase in observed distress scores. The model's bounded "
  "benefit means the expected effect is small; a trial powered for a large "
  "effect would fail for the wrong reason."),
 ("P5. Preserved responsiveness. ",
  "If the mechanism operates, the intervention arm should retain measurable "
  "responsiveness longer. This is the outcome that matters to families and the "
  "one the hypothesis exists to serve, and it should be reported whatever the "
  "dose result."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(5)
    r = p.add_run(tag)
    r.bold = True
    r.font.size = Pt(11)
    r = p.add_run(body)
    r.font.size = Pt(11)
Pp("P3 is the most direct early test of the proposed auditory–autonomic "
   "link. It does not test the rest of the chain to opioid escalation, which P2 "
   "and P4 address, but it is non-invasive, requires no protocol deviation, and "
   "would be informative whichever way it resolved.", indent=True)
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.5
p.paragraph_format.space_before = Pt(5)
r = p.add_run("Feasibility of P3. ")
r.bold = True
r.font.size = Pt(11)
r = p.add_run(
    "High-frequency heart rate variability is hard to measure in actively dying "
    "patients, and a protocol has to plan for that rather than discover it. "
    "Atrial fibrillation and frequent ectopy make short-term HRV uninterpretable, "
    "so patients in atrial fibrillation are excluded and ectopic beats are "
    "corrected or excluded by rules fixed in advance %s. High-frequency power "
    "tracks respiratory sinus arrhythmia, so changes in breathing rate or pattern "
    "— including periodic and terminal breathing — move it independently "
    "of vagal tone; respiration must be recorded alongside the ECG and "
    "high-frequency indices interpreted only when breathing lies within the "
    "0.15–0.40 Hz band %s. Anticholinergics given for respiratory "
    "secretions, and opioids and sedatives themselves, alter cardiac vagal "
    "activity, so every dose is time-stamped and blocks adjacent to a bolus are "
    "excluded. Signal quality should come from ECG rather than "
    "photoplethysmography, at a sampling rate of at least 250 Hz %s, with the "
    "proportion of edited beats reported per segment and segments above a "
    "preset threshold discarded. Because the autonomic state of a dying patient "
    "drifts over hours, the comparison should be within patient: alternating "
    "short blocks of structured verbal input and a matched acoustic control, "
    "with the change in log high-frequency power or RMSSD as the outcome. The "
    "first study should report feasibility itself — the fraction of eligible "
    "patients and of recorded blocks that yield analysable data — because if "
    "that fraction is small, P3 cannot be tested in this population by this "
    "method."
    % (C("taskforce1996"), C("laborde2017"), C("taskforce1996")))
r.font.size = Pt(11)

# ================================================================= SECTION 8
H("8. How this proposal could cause harm")
Pp("A hypothesis whose practical implication is a smaller opioid dose in dying "
   "patients carries a specific and serious risk, and enumerating it is part of "
   "the proposal rather than a caveat appended to it.")
for lead, body in [
 ("Under-treatment of pain. ",
  "The principal danger is that the framework is read as a reason to withhold "
  "or delay analgesia in a patient who cannot report pain. Nothing here "
  "supports that. The proposal is adjunctive: the opioid is titrated to comfort "
  "exactly as it would otherwise be, and any reduction in requirement is an "
  "observed outcome, never a target to be pursued."),
 ("Misattribution of distress. ",
  "If a patient appears calmer because arousal has been damped while "
  "nociception is unchanged, observed distress scores could fall without the "
  "underlying pain changing. This is the failure mode most likely to cause "
  "harm and is the reason P4 specifies distress scores as a safety endpoint "
  "rather than an efficacy one."),
 ("Burden on the dying and on families. ",
  "Continuous verbal intervention at the bedside is not neutral. It occupies "
  "the room, may displace silence that families need, and may be experienced "
  "as intrusive by patients who cannot decline it. Consent cannot be obtained "
  "from an unresponsive patient, and advance preference should govern."),
 ("False reassurance. ",
  "Presenting an unproven adjunct as a way to preserve lucidity risks families "
  "attributing an unavoidable decline to insufficient effort, or declining "
  "adequate analgesia in hope of more time. The bounded benefit in Section 5 "
  "should be communicated as a bound, not as an expectation."),
 ("Displacement of established care. ",
  "Music and hypnosis interventions in palliative care have been "
  "systematically reviewed: they appear feasible and acceptable, with a "
  "moderate reduction in pain across a small number of randomised trials, but "
  "the evidence is too limited to compare interventions or establish effects "
  "on most other outcomes %s. A new adjunct with weaker evidence should not "
  "compete for the same resources." % C("bissonnette2024")),
]:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(5)
    r = p.add_run(lead)
    r.bold = True
    r.font.size = Pt(11)
    r = p.add_run(body)
    r.font.size = Pt(11)

# ================================================================= SECTION 9
H("9. Relation to existing practice")
Pp("Soothing speech, familiar voices and music are already part of good hospice "
   "practice, hypnotherapy has been delivered to palliative patients for "
   "anxiety %s, and music and hypnosis interventions in palliative care have "
   "been evaluated for pain, anxiety, sleep and well-being %s. The claim here "
   "is narrower than 'sound helps'. It is that a "
   "specific, measurable autonomic pathway mediates the effect; that the "
   "relevant outcome is the trajectory of opioid requirement rather than a "
   "comfort score; and that the mechanism predicts a bounded, quantifiable "
   "benefit which can be tested and found absent. Existing practice is "
   "compatible with the hypothesis but does not test it: to our knowledge, it "
   "has not been paired with autonomic measurement or opioid dose trajectory "
   "as endpoints." % (C("plaskota2012"), C("bissonnette2024")))

H("10. Ethical statement")
Pp("This is a conceptual hypothesis paper. No patient was studied, no data were "
   "collected, and no intervention described here has been tested in any "
   "population. Nothing in it constitutes clinical advice or a treatment "
   "recommendation, and nothing in it should delay, reduce or substitute for "
   "established palliative care, including adequate analgesia. Decisions about "
   "opioid dosing in dying patients are clinical judgements made at the bedside "
   "by those responsible for the patient. Any prospective evaluation would "
   "require research ethics approval with particular attention to consent by "
   "proxy, to the right to decline, and to safety monitoring for under-treated "
   "pain.")

# ============================================================== DECLARATIONS
H("Declarations")
Pp("Funding. This work received no external funding.")
Pp("Competing interests. The author declares no competing interests.")
Pp("Data availability. The model implementation, parameter set and figure "
   "generators are openly available at %s and archived at https://doi.org/%s. "
   "Running escalation_model.py followed by make_figures.py reproduces every "
   "number and figure. This manuscript is deposited at "
   "https://doi.org/10.5281/zenodo.22860432. Both are concept DOIs and resolve "
   "to the current version. The equations, initial conditions and parameter "
   "values needed to reproduce the model without the code are given in Section 5 "
   "and Tables 1–2. No patient data were generated or analysed."
   % ("https://github.com/sandlerleon/auditory-cardiac-model", "10.5281/zenodo.22860430"))

H("Declaration of generative AI and AI-assisted technologies in the manuscript "
  "preparation process", size=11)
Pp("During the preparation of this work the author used Claude (Anthropic) in "
   "order to assist with literature synthesis, implementation of the "
   "illustrative model, and drafting and editing of the text. After using this "
   "tool, the author reviewed and edited the content as needed, verified every "
   "cited reference against its source, and takes full responsibility for the "
   "content of the published article. No generative AI tool is listed as an "
   "author, and none was used to produce or alter data.")

# ================================================================ REFERENCES
H("References")


def _end(s):
    """Terminate with a full stop unless the text already ends in punctuation."""
    return s if s[-1] in ".?!" else s + "."


ARTICLE_NO = {"blundon2020": "10336"}              # Crossref article-number


def fmt_ref(v, tag=None):
    au = _end(", ".join(v["authors"][:6]) + (", et al." if len(v["authors"]) > 6 else ""))
    ti = _end(v["title"])
    if v.get("url") and not v.get("doi"):          # guideline without a DOI
        return "%s %s %s; %d. %s" % (au, ti, v["journal"], v["year"], v["url"])
    if v.get("cochrane"):                          # Cochrane: year;(issue):article number
        loc = ";" + v["cochrane"]
    elif v.get("volume"):
        pg = v.get("page") or ARTICLE_NO.get(tag)
        loc = ";" + v["volume"] + ((":%s" % pg) if pg else "")
    else:                                          # online ahead of print
        loc = ". Published online; article %s" % v["page"] if v.get("page") else ""
    return "%s %s %s %d%s. https://doi.org/%s" % (au, ti, v["journal"], v["year"], loc,
                                                  v["doi"])


for i, tag in enumerate(ORDER, 1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)
    r = p.add_run("[%d] %s" % (i, fmt_ref(REFS[tag], tag)))
    r.font.size = Pt(9.5)

if not os.path.isdir(OUT):
    os.makedirs(OUT)
path = os.path.join(OUT, "Auditory_Cardiac_Channel_MedicalHypotheses.docx")
doc.save(path)

words = sum(len(p.text.split()) for p in doc.paragraphs)
print("abstract words : %d" % len(ABSTRACT.split()))
print("total words    : %d" % words)
print("references     : %d (cited, in order of appearance)" % len(ORDER))
print("unused refs    : %s" % (sorted(set(REFS) - set(ORDER)) or "none"))
print("saved          : %s" % path)
